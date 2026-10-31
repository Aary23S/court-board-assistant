"""
Word (DOCX) Renderer for Court-Approved Roznama Proceedings.
Produces professionally formatted, LibreOffice Writer-compatible documents
matching the visual 3-column table structure, repeated headings, and dotted separators.
"""
from pathlib import Path
from typing import Union, List, Optional
import docx
from docx.shared import Pt, Inches, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from ..domain.models import RoznamaRecord, RoznamaHeader
from .presentation_model import RoznamaPresentationModel

class RoznamaWordRenderer:
    """
    Renders RoznamaPresentationModel into an editable Word (.docx) document.
    """
    def __init__(self):
        pass

    def _set_cell_border(self, cell, **kwargs):
        """
        Sets borders on a table cell.
        kwargs: top, bottom, left, right -> {'sz': val, 'val': val, 'color': val}
        """
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = tcPr.first_child_found_in("w:tcBorders")
        if tcBorders is None:
            tcBorders = OxmlElement('w:tcBorders')
            tcPr.append(tcBorders)
        for edge, border_args in kwargs.items():
            edge_el = tcBorders.find(qn(f'w:{edge}'))
            if edge_el is None:
                edge_el = OxmlElement(f'w:{edge}')
                tcBorders.append(edge_el)
            for key, val in border_args.items():
                edge_el.set(qn(f'w:{key}'), str(val))

    def _set_cell_margins(self, cell, top=100, bottom=100, left=150, right=150):
        """Sets internal padding (in twips) for a cell."""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def render(
        self,
        presentation_model: Union[RoznamaPresentationModel, List[RoznamaRecord]],
        output_path: Union[str, Path],
        header: Optional[RoznamaHeader] = None
    ) -> Path:
        """
        Renders the presentation model into a Word (.docx) document at output_path.
        """
        if not isinstance(presentation_model, RoznamaPresentationModel):
            presentation_model = RoznamaPresentationModel.from_state_records(
                transformed_records=presentation_model,
                source_records=[],
                human_edits={},
                header_case_number=header.case_number if header else None
            )

        out_path = Path(output_path).resolve()
        out_path.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()

        # Page Setup: A4 Portrait, Margins: Top 15mm, Bottom 15mm, Left 20mm, Right 10mm
        section = doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(15)
        section.bottom_margin = Mm(15)
        section.left_margin = Mm(20)
        section.right_margin = Mm(10)

        # Usable width = 210 - 20 - 10 = 180mm (~7.086 inches / ~510pt)
        col1_w = Mm(33.5)   # ~18.6%
        col2_w = Mm(106.5)  # ~59.1%
        col3_w = Mm(40.0)   # ~22.3%

        font_name = "Lohit Devanagari"

        total_items = len(presentation_model.items)

        for idx, item in enumerate(presentation_model.items):
            # 1. Repeated Case Heading
            p_head = doc.add_paragraph()
            p_head.paragraph_format.space_before = Pt(6)
            p_head.paragraph_format.space_after = Pt(3)
            p_head.paragraph_format.keep_with_next = True
            run_head = p_head.add_run(item.case_heading)
            run_head.font.name = font_name
            run_head.font.size = Pt(10.5)
            run_head.font.bold = True
            run_head.font.color.rgb = RGBColor(0, 0, 0)

            # 2. Proceedings Table (3 Columns)
            table = doc.add_table(rows=2, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            # Prevent rows from splitting awkwardly across pages
            for row in table.rows:
                trPr = row._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            # Repeat header row if table crosses page
            header_tr = table.rows[0]._tr.get_or_add_trPr()
            header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            # Setup Header Cells
            header_titles = ["आजची दिनांक", "दैनंदिन कामकाज तपशील", "पुढे/निकाल तारीख"]
            border_spec = {'sz': 4, 'val': 'single', 'color': '000000'}

            for c_idx, title in enumerate(header_titles):
                cell = table.cell(0, c_idx)
                self._set_cell_border(cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
                self._set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAFAFA"/>')
                cell._tc.get_or_add_tcPr().append(shading_elm)

                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(title)
                run.font.name = font_name
                run.font.size = Pt(10)
                run.font.bold = True

            # Setup Data Cells
            # Col 1: Date
            cell_date = table.cell(1, 0)
            self._set_cell_border(cell_date, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
            self._set_cell_margins(cell_date, top=100, bottom=100, left=100, right=100)
            p_d = cell_date.paragraphs[0]
            p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_d = p_d.add_run(item.hearing_date)
            run_d.font.name = font_name
            run_d.font.size = Pt(10)

            # Col 2: Business (Line 1) + Stage (Line 2) + Judge designation
            cell_detail = table.cell(1, 1)
            self._set_cell_border(cell_detail, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
            self._set_cell_margins(cell_detail, top=100, bottom=100, left=100, right=100)

            p_b = cell_detail.paragraphs[0]
            p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_b.paragraph_format.space_before = Pt(0)
            p_b.paragraph_format.space_after = Pt(2)
            run_b = p_b.add_run(item.business_text)
            run_b.font.name = font_name
            run_b.font.size = Pt(9.5)

            if item.stage_text:
                p_s = cell_detail.add_paragraph()
                p_s.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p_s.paragraph_format.space_before = Pt(2)
                p_s.paragraph_format.space_after = Pt(4)
                run_s = p_s.add_run(item.stage_text)
                run_s.font.name = font_name
                run_s.font.size = Pt(9.5)

            if item.judge_designation:
                p_j = cell_detail.add_paragraph()
                p_j.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p_j.paragraph_format.space_before = Pt(6)
                p_j.paragraph_format.space_after = Pt(2)
                run_j = p_j.add_run(item.judge_designation)
                run_j.font.name = font_name
                run_j.font.size = Pt(9)
                run_j.font.bold = True

            # Col 3: Next Date
            cell_next = table.cell(1, 2)
            self._set_cell_border(cell_next, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
            self._set_cell_margins(cell_next, top=100, bottom=100, left=100, right=100)
            p_n = cell_next.paragraphs[0]
            p_n.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_n = p_n.add_run(item.next_date)
            run_n.font.name = font_name
            run_n.font.size = Pt(10)

            # Apply explicit column widths to cells
            widths = [col1_w, col2_w, col3_w]
            for row in table.rows:
                for c_idx, w in enumerate(widths):
                    row.cells[c_idx].width = w

            # 3. Dotted Separator between consecutive hearings (omit after final)
            if idx < total_items - 1:
                p_sep = doc.add_paragraph()
                p_sep.paragraph_format.space_before = Pt(4)
                p_sep.paragraph_format.space_after = Pt(4)
                # Dotted bottom border on paragraph
                pPr = p_sep._p.get_or_add_pPr()
                pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="dotted" w:sz="8" w:space="1" w:color="666666"/></w:pBdr>')
                pPr.append(pBdr)

        doc.save(str(out_path))
        return out_path

